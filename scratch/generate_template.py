import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image
import os

# Create a new workbook and select the active worksheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Plantilla Cronograma"

# Institutional Colors
UDEC_YELLOW = "F1C40F"
UDEC_DARK = "2C3E50"

# Set up column widths
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 15
ws.column_dimensions['D'].width = 15
ws.column_dimensions['E'].width = 15
ws.column_dimensions['F'].width = 15
ws.column_dimensions['G'].width = 15
ws.column_dimensions['H'].width = 25

# 1. Skip Image
logo_path = 'public/logo-udec.png'

# 2. Institutional Header
ws.merge_cells('C1:H1')
ws['C1'] = "UNIVERSIDAD DE CARTAGENA"
ws['C1'].font = Font(bold=True, size=16, color="FFFFFF")
ws['C1'].alignment = Alignment(horizontal="center", vertical="center")
ws['C1'].fill = PatternFill(start_color=UDEC_DARK, end_color=UDEC_DARK, fill_type="solid")

ws.merge_cells('C2:H2')
ws['C2'] = "FACULTAD DE CIENCIAS ECONÓMICAS"
ws['C2'].font = Font(bold=True, size=14)
ws['C2'].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells('C3:H3')
ws['C3'] = "DEPARTAMENTO DE POSGRADOS Y EDUCACIÓN CONTINUA"
ws['C3'].font = Font(bold=True, size=12)
ws['C3'].alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells('C4:H4')
ws['C4'] = "PLANTILLA ESTANDARIZADA DE CRONOGRAMA ACADÉMICO - IPAS"
ws['C4'].font = Font(bold=True, size=12, color="FFFFFF")
ws['C4'].alignment = Alignment(horizontal="center", vertical="center")
ws['C4'].fill = PatternFill(start_color="808080", end_color="808080", fill_type="solid")

# Row heights for header
ws.row_dimensions[1].height = 25
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 20
ws.row_dimensions[4].height = 20

# 3. Basic Program Info
ws['A6'] = "Programa:"
ws['A6'].font = Font(bold=True)
ws['B6'] = "[Nombre del Posgrado]"
ws['C6'] = "Cohorte:"
ws['C6'].font = Font(bold=True)
ws['D6'] = "[Nro]"
ws['E6'] = "Semestre:"
ws['E6'].font = Font(bold=True)
ws['F6'] = "I - 2027"

# 4. Table Headers
headers = [
    "Módulo / Asignatura",
    "Docente(s)",
    "Fecha de Inicio",
    "Fecha de Fin",
    "Horario",
    "Créditos",
    "Horas",
    "Observaciones / Requisitos Esp."
]

header_font = Font(bold=True, color="000000")
header_fill = PatternFill(start_color=UDEC_YELLOW, end_color=UDEC_YELLOW, fill_type="solid")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=8, column=col_num)
    cell.value = header
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

ws.row_dimensions[8].height = 30

# 5. Empty Rows for Data Entry
for row in range(9, 29):
    for col in range(1, 9):
        cell = ws.cell(row=row, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center", wrap_text=True)

# Save the workbook
output_path = 'Plantilla_Cronograma_IPAS.xlsx'
wb.save(output_path)
print(f"Plantilla generada exitosamente en {output_path}")
